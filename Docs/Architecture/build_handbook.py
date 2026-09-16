"""Build offline HTML and editable DOCX from the adjacent Markdown sources.

Dependencies: markdown, python-docx (which supplies lxml).
Use a project/document virtual environment, then:
    python Docs/Architecture/build_handbook.py

For this workspace's temporary packages:
    <python> -X utf8 Docs/Architecture/build_handbook.py \
        --deps Saved/ArchitectureWork/python-packages

No network requests, Unreal changes, or source-document edits are performed.
Markdown remains authoritative. DOCX renders diagrams as readable text flows;
HTML contains offline SVG diagrams and the original Mermaid source.
"""

from pathlib import Path
import argparse
import html
import re
import sys
import textwrap
from urllib.parse import unquote

parser = argparse.ArgumentParser()
parser.add_argument('--deps', type=Path)
args = parser.parse_args()
if args.deps:
    sys.path.insert(0, str(args.deps.resolve()))

import markdown
from lxml import etree
from lxml import html as lh
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
FILES = [
    'README.md', '01-Foundation-and-Code-Setup.md',
    '02-Gameplay-System-Contracts.md', '03-Authoring-and-Team-Workflows.md',
    '04-Implementation-and-Validation.md',
]


def svg_diagram(nodes, edges, width, height, title):
    """Render project-authored diagrams as accessible, self-contained vectors."""
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" role="img" aria-label="{html.escape(title)}">',
           '<defs><marker id="ARROW" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#57708a"/></marker></defs>']
    for edge in edges:
        a, b = edge[:2]
        ax, ay, aw, ah, _ = nodes[a]
        bx, by, bw, bh, _ = nodes[b]
        acx, acy, bcx, bcy = ax + aw/2, ay + ah/2, bx + bw/2, by + bh/2
        dx, dy = bcx-acx, bcy-acy
        sa = min(aw/2/abs(dx) if dx else 1e9, ah/2/abs(dy) if dy else 1e9)
        sb = min(bw/2/abs(dx) if dx else 1e9, bh/2/abs(dy) if dy else 1e9)
        x1, y1, x2, y2 = acx+dx*sa, acy+dy*sa, bcx-dx*sb, bcy-dy*sb
        out.append(f'<path d="M{x1},{y1} L{x2},{y2}" fill="none" stroke="#57708a" stroke-width="2" marker-end="url(#ARROW)"/>')
        if len(edge) > 2:
            out.append(f'<text x="{(x1+x2)/2+8}" y="{(y1+y2)/2-7}" font-family="Segoe UI,Arial,sans-serif" font-size="13" fill="#334b65">{html.escape(edge[2])}</text>')
    for x, y, w, h, label in nodes.values():
        out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="9" fill="#eef5fa" stroke="#83a4bc"/>')
        lines = label.split('\n')
        for i, line in enumerate(lines):
            ty = y+h/2+(i-(len(lines)-1)/2)*19+5
            out.append(f'<text x="{x+w/2}" y="{ty}" text-anchor="middle" font-family="Segoe UI,Arial,sans-serif" font-size="14" fill="#17364c">{html.escape(line)}</text>')
    out.append('</svg>')
    return ''.join(out)


def diagram_for(source, index):
    if 'sequenceDiagram' in source:
        labels = ['UI submits purchase', 'Prepare required assets', 'Revalidate live request', 'Commit debit + rank', 'Publish result + revision', 'Queue coherent save']
        plain = '\n'.join(f'{i+1}. {v}' for i, v in enumerate(labels))
        nodes = {str(i): (30+(i%3)*290, 25+(i//3)*140, 245, 70, v) for i, v in enumerate(labels)}
        # A readable two-row sequence, without a diagonal return connector.
        nodes['3'], nodes['5'] = (610,165,245,70,labels[3]), (30,165,245,70,labels[5])
        edges = [('0','1'),('1','2'),('2','3'),('3','4'),('4','5')]
        visual = svg_diagram(nodes, edges, 890, 265, 'Purchase transaction sequence')
    elif 'SOTM1' in source:
        nodes = {'app':(335,15,220,65,'SOTM1\nApplication bootstrap'), 'game':(170,145,220,65,'SOTMGameplay\nRules and state'), 'view':(505,145,220,65,'SOTMPresentation\nUI and feedback'), 'core':(335,285,220,65,'SOTMCore\nShared contracts'), 'edit':(770,15,200,65,'SOTMEditor\nAuthoring tools')}
        edges = [('app','game'),('app','view'),('view','game'),('game','core'),('view','core'),('edit','view'),('edit','game')]
        visual = svg_diagram(nodes, edges, 995, 380, 'Module compile dependencies')
        plain = 'Compile dependencies:\nSOTM1 → SOTMGameplay and SOTMPresentation\nSOTMPresentation → SOTMGameplay and SOTMCore\nSOTMGameplay → SOTMCore\nSOTMEditor → SOTMGameplay and SOTMPresentation\nRuntime gameplay never depends on presentation or editor modules.'
    elif 'Stamina' in source:
        nodes = {'in':(20,20,180,65,'Input intent'), 'fact':(20,135,180,75,'Stamina, stance,\neffects, restrictions'), 'pol':(265,75,180,70,'Movement policy'), 'move':(510,75,180,70,'Character Movement'), 'actual':(755,75,180,70,'Actual velocity\nand movement mode'), 'anim':(755,205,180,70,'Animation\nand camera'), 'noise':(265,205,180,70,'Gameplay noise')}
        visual = svg_diagram(nodes, [('in','pol'),('fact','pol'),('pol','move'),('move','actual'),('actual','anim'),('pol','noise')], 960,300,'Movement ownership')
        plain = 'Input intent + stamina/stance/effects/restrictions\n    → Movement policy → Character Movement\n    → Actual velocity/movement mode → Animation and camera\nMovement policy also drives the gameplay noise model.'
    elif 'Death accepted' in source:
        nodes = {'play':(20,25,190,65,'Playing'), 'accept':(265,25,220,65,'Death accepted\nConsume one life'), 'pres':(545,25,220,65,'Kill presentation\nor fallback'), 'screen':(545,160,220,65,'Death screen'), 'restore':(265,160,220,65,'Lives > 0:\nrestore checkpoint'), 'over':(545,285,220,65,'Lives = 0:\nGame Over'), 'new':(265,285,220,65,'New run or menu')}
        visual = svg_diagram(nodes,[('play','accept'),('accept','pres'),('pres','screen'),('screen','restore'),('restore','play'),('screen','over'),('over','new')],800,380,'Five lives and retry flow')
        plain = 'Playing → Death accepted (consume one life)\n    → Kill presentation or fallback → Death screen\nIf lives > 0: restore checkpoint → Playing\nIf lives = 0: Game Over → New run or menu'
    else:
        labels = ['Boot and profile', 'Frontend', 'Select chapter/profile', 'Load assets and map', 'Restore world records', 'Possess/rebuild abilities', 'Refresh views / ready', 'Playing', 'Travel or retry']
        positions = [(20,20),(310,20),(600,20),(600,145),(310,145),(20,145),(20,270),(310,270),(600,270)]
        nodes = {str(i):(*positions[i],250,65,v) for i,v in enumerate(labels)}
        visual = svg_diagram(nodes, [(str(i),str(i+1)) for i in range(8)]+[('8','3')],880,365,'Boot and chapter lifecycle')
        plain = 'Boot/profile → Frontend → Select chapter/build profile\n    → Load assets/map → Register and restore world records\n    → Possess player/rebuild abilities → Refresh views/readiness\n    → Playing\nTravel or checkpoint retry repeats required preparation/restoration.'
    # Unique SVG markers when several diagrams share one HTML document.
    visual = visual.replace('ARROW', f'arrow-{index}')
    return visual, plain


sections, headings, diagram_texts = [], [], {}
word_count = 0
for number, name in enumerate(FILES):
    source = (ROOT/name).read_text(encoding='utf-8')
    word_count += len(source.split())
    rendered = markdown.markdown(source, extensions=['extra', 'sane_lists', 'toc'])
    tree = lh.fragment_fromstring(rendered, create_parent='section')
    tree.set('id', f'part-{number}')
    for heading in tree.xpath('.//h1|.//h2|.//h3|.//h4'):
        hid = f'p{number}-' + heading.get('id', 'heading')
        heading.set('id', hid)
        headings.append((int(heading.tag[1]), hid, heading.text_content()))
    for link in tree.xpath('.//a[@href]'):
        href = link.get('href')
        base, _, anchor = href.partition('#')
        if base in FILES:
            idx = FILES.index(base)
            link.set('href', f'#p{idx}-{anchor}' if anchor else f'#part-{idx}')
        elif not base and anchor:
            link.set('href', f'#p{number}-{anchor}')
    for pre in tree.xpath('.//pre[code[contains(@class,"language-mermaid")]]'):
        code = pre[0].text or ''
        index = len(diagram_texts)+1
        svg, plain = diagram_for(code,index)
        figure = lh.fragment_fromstring(f'<figure class="diagram" data-diagram="{index}">{svg}<details><summary>Diagram source (Mermaid)</summary><pre><code>{html.escape(code)}</code></pre></details></figure>')
        diagram_texts[str(index)] = plain
        pre.getparent().replace(pre,figure)
    sections.append(tree)

toc = ''.join(f'<li class="depth-{level}"><a href="#{hid}">{html.escape(label)}</a></li>' for level,hid,label in headings if level<=2)
content = '\n'.join(lh.tostring(s, encoding='unicode') for s in sections)
css = '''
:root{--nav:#102c40;--ink:#193044;--accent:#126d85;--line:#d7e2e9;--muted:#546779}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:24px}
body{margin:0;background:#f4f7fa;color:var(--ink);font:16px/1.65 "Segoe UI",Arial,sans-serif}
nav{position:fixed;inset:0 auto 0 0;width:285px;background:var(--nav);color:#e4eef5;padding:26px 22px;overflow:auto}
.brand{font-size:12px;font-weight:700;letter-spacing:2px;color:#79c5d6}.navtitle{font-size:24px;font-weight:650;line-height:1.25;margin:10px 0 18px}
nav input{width:100%;padding:10px;border:1px solid #5c798c;border-radius:5px;background:#1c3a4e;color:white;margin-bottom:14px}
nav input::placeholder{color:#c8d8e1}nav ul{list-style:none;padding:0}nav li{line-height:1.35;margin:10px 0}nav a{color:#dbe9f1;text-decoration:none;font-size:13px}nav a:hover{color:#8dd9e7}
nav .depth-1{font-weight:700;margin-top:22px}nav .depth-2{padding-left:12px;border-left:1px solid #486576}
main{margin-left:285px;padding:48px 5vw 90px;max-width:1570px}article{background:white;border:1px solid var(--line);border-radius:10px;padding:40px 48px;box-shadow:0 8px 25px #132c4009}
.masthead{margin-bottom:28px}.eyebrow{letter-spacing:2px;text-transform:uppercase;font-size:12px;font-weight:700;color:var(--accent)}.masthead p{color:var(--muted);margin:8px 0}
button{border:1px solid #24748a;color:#17687e;background:white;border-radius:5px;padding:8px 13px;cursor:pointer;font:inherit;font-size:13px}
h1{font-size:35px;line-height:1.2;letter-spacing:-.6px;margin:0 0 20px;color:#123449}h2{font-size:25px;line-height:1.3;margin:44px 0 18px;padding-bottom:9px;border-bottom:2px solid #e2edf2;color:#174d65}h3{font-size:19px;margin:30px 0 12px;color:#23627b}
section+section{margin-top:65px;padding-top:50px;border-top:5px solid #1c7188}p{margin:13px 0}a{color:#116d89;text-decoration-thickness:1px;text-underline-offset:2px}strong{color:#163c51}
table{width:100%;border-collapse:collapse;font-size:13px;line-height:1.5;margin:22px 0;table-layout:auto}th{background:#164760;color:white;text-align:left;font-weight:650}td,th{padding:10px 11px;border:1px solid #d7e2e9;vertical-align:top;overflow-wrap:anywhere}tr:nth-child(even) td{background:#f3f7fa}
pre{background:#f0f5f8;border:1px solid #d4e1e8;border-left:4px solid #43879e;padding:16px;overflow:auto;line-height:1.5;border-radius:4px;font-size:12px}code{font-family:Consolas,"Cascadia Code",monospace;font-size:.88em}p code,li code,td code{background:#eef3f6;border-radius:3px;padding:1px 4px;overflow-wrap:anywhere}
li{margin:8px 0}figure{margin:28px 0;padding:15px;border:1px solid #dbe6ed;border-radius:8px;background:#fcfeff}svg{display:block;width:100%;height:auto}details{font-size:12px;color:var(--muted)}summary{cursor:pointer}.footer{margin-top:30px;font-size:12px;color:var(--muted)}
@media(max-width:1100px){nav{width:240px}main{margin-left:240px;padding:25px}article{padding:30px}h1{font-size:30px}}
@media(max-width:760px){nav{position:relative;width:auto;max-height:340px}main{margin:0;padding:15px}article{padding:22px 16px}table{font-size:11px}td,th{padding:7px}h1{font-size:27px}}
@media print{@page{size:A4;margin:16mm}body{background:white;font-size:10pt;line-height:1.45}nav,.masthead,button,details{display:none}main{margin:0;padding:0;max-width:none}article{padding:0;border:0;box-shadow:none}h1{font-size:24pt}h2{font-size:17pt;break-after:avoid}h3{font-size:13pt;break-after:avoid}section+section{break-before:page;border:0;padding-top:0}table{font-size:8pt}thead{display:table-header-group}tr,figure{break-inside:avoid}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:7pt}a{color:inherit}a[href^="http"]{overflow-wrap:anywhere}.footer{display:none}}
'''
page = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SOTM — Unreal Game Architecture Handbook</title><style>{css}</style></head><body>
<nav aria-label="Handbook contents"><div class="brand">SOTM / ENGINEERING</div><div class="navtitle">Game Architecture<br>Handbook</div><input id="toc-filter" aria-label="Filter table of contents" placeholder="Filter contents…"><ul id="toc">{toc}</ul></nav>
<main><div class="masthead"><div class="eyebrow">Architecture · Authoring · Validation</div><p>Revision 1 / 9 September 2026 / Proposed baseline: Unreal Engine 5.8.2</p><button type="button" onclick="window.print()">Print / Save as PDF</button> <span style="font-size:13px;color:#546779">Use Ctrl+F to search the full handbook.</span></div><article>{content}</article><div class="footer">Generated from the accompanying Markdown. Offline diagrams and navigation; no external scripts, fonts, or analytics.</div></main>
<script>document.getElementById('toc-filter').addEventListener('input',function(){{const q=this.value.toLowerCase();document.querySelectorAll('#toc li').forEach(li=>li.hidden=!li.textContent.toLowerCase().includes(q));}});</script></body></html>'''
(ROOT/'Architecture-Handbook.html').write_text(page,encoding='utf-8')

# Editable Word rendering: native headings, tables, hyperlinks and bookmarks.
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.65)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for level,size in [(1,24),(2,17),(3,13),(4,11)]:
    style=doc.styles[f'Heading {level}']
    style.font.name,style.font.size,style.font.color.rgb='Calibri',Pt(size),RGBColor.from_string('174D65')
    style.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]
header.text='SOTM  /  GAME ARCHITECTURE HANDBOOK'
header.style='Caption'
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Revision 1  •  9 September 2026  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

bookmark_ids={}
for idx,(_,hid,_) in enumerate(headings,1):
    bookmark_ids[hid]=f'h{idx}'
for idx in range(len(FILES)):
    first=next(hid for _,hid,_ in headings if hid.startswith(f'p{idx}-'))
    bookmark_ids[f'part-{idx}']=bookmark_ids[first]


def inline(paragraph, element, bold=False, italic=False, code=False):
    def add(text,b=bold,i=italic,c=code):
        if not text:return
        run=paragraph.add_run(text);run.bold=b;run.italic=i
        if c:run.font.name='Consolas';run.font.size=Pt(8)
    add(element.text)
    for child in element:
        tag=child.tag
        if tag=='a':
            target=child.get('href','');label=child.text_content()
            link=OxmlElement('w:hyperlink')
            if target.startswith('#') and target[1:] in bookmark_ids:
                link.set(qn('w:anchor'),bookmark_ids[target[1:]])
            else:
                link.set(qn('r:id'),paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'126D85');props.append(color);run.append(props)
            text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)
        elif tag=='br':paragraph.add_run().add_break()
        else:inline(paragraph,child,bold or tag in ('strong','b'),italic or tag in ('em','i'),code or tag=='code')
        add(child.tail)


def blocks(parent, container, first_section=False):
    for el in parent:
        tag=el.tag
        if tag in ('h1','h2','h3','h4'):
            level=int(tag[1]);p=container.add_paragraph(style=f'Heading {level}')
            if level==1 and not first_section:p.paragraph_format.page_break_before=True
            inline(p,el)
            key=bookmark_ids.get(el.get('id'))
            if key:
                start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),key[1:]);start.set(qn('w:name'),key)
                end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),key[1:]);p._p.insert(0,start);p._p.append(end)
        elif tag=='p':inline(container.add_paragraph(),el)
        elif tag in ('ul','ol'):
            for item in el:
                p=container.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number')
                if len(item) and item[0].tag=='p':
                    inline(p,item[0]);blocks(item,container) if len(item)>1 else None
                else:inline(p,item)
        elif tag=='table':
            rows=el.xpath('./thead/tr|./tbody/tr|./tr')
            columns=max(len(r) for r in rows)
            table=container.add_table(rows=0,cols=columns);table.style='Light Shading Accent 1'
            for i,row in enumerate(rows):
                cells=table.add_row().cells
                for j,cell in enumerate(row):
                    p=cells[j].paragraphs[0];inline(p,cell)
                    for r in p.runs:r.font.size=Pt(8)
                trpr=table.rows[-1]._tr.get_or_add_trPr()
                keep=OxmlElement('w:cantSplit');trpr.append(keep)
                if i==0:
                    repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
                    for c in cells:
                        for r in c.paragraphs[0].runs:r.bold=True
            container.add_paragraph()
        elif tag in ('pre','figure'):
            value=diagram_texts[el.get('data-diagram')] if tag=='figure' else el.text_content()
            for line in value.strip().splitlines():
                p=container.add_paragraph();p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1
                r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(7.5)
            container.add_paragraph()
        elif tag=='blockquote':
            p=container.add_paragraph(style='Quote');p.add_run(el.text_content())
        elif tag not in ('hr',):blocks(el,container)

for i,section in enumerate(sections):
    blocks(section,doc,first_section=i==0)
doc.core_properties.title='SOTM Unreal Game Architecture Handbook'
doc.core_properties.subject='Architecture setup, gameplay contracts, authoring workflows, implementation and validation'
doc.core_properties.author='SOTM project documentation'
doc.core_properties.keywords='Unreal Engine, horror, architecture, data driven, gameplay, authoring'
doc.save(ROOT/'Architecture-Handbook.docx')

# Structural validation of generated output and all source-local links.
tree=lh.fromstring(page)
ids=tree.xpath('//@id')
assert len(ids)==len(set(ids)), 'Duplicate HTML identifiers'
for anchor in tree.xpath('//a[starts-with(@href,"#")]/@href'):
    assert anchor[1:] in ids, f'Broken HTML anchor: {anchor}'
for name in FILES:
    text=(ROOT/name).read_text(encoding='utf-8')
    assert text.count('```')%2==0,f'Unclosed code fence: {name}'
    for url in re.findall(r'\]\(([^)]+)\)',text):
        if url.startswith(('https://','http://','#')):continue
        path=url.partition('#')[0]
        assert (ROOT/unquote(path)).exists(),f'Broken local link in {name}: {url}'
assert len(tree.xpath('//svg'))==5,'Expected five rendered architecture diagrams'
reopened=Document(ROOT/'Architecture-Handbook.docx')
assert len(reopened.tables)==len(tree.xpath('//article//table')),'Word table loss'
assert len([p for p in reopened.paragraphs if p.style.name.startswith('Heading')])==len(headings),'Word heading loss'
print(f'Built HTML and DOCX: {word_count:,} source words; {len(headings)} headings; {len(reopened.tables)} tables; 5 diagrams.')
print('Passed local-link, anchor, fence, Word reopen, heading and table checks.')
